"""报告渲染与落盘（F15）：把报告结构渲染成 json / md / html 三份产物。

职责边界（刻意收窄）：
- 本模块**只做渲染**，不做任何业务计算——通过率 / 覆盖率 / 趋势 / 结论一律来自传入的
  `report` 字典（由 `engine/report.py` 计算）。这样渲染器可被单独替换（换皮肤/换模板）
  而不影响结论口径。
- HTML 为**自包含浅色页面**（内联样式、零外部依赖），便于直接打开与分享给非技术同学。
- 所有外部文本（用例标题、备注、源码路径）一律转义，防止标题里的 `<` / `&` 破坏页面结构。

产物目录（与代码只读区物理分离）：
    outputs/<project_id>/REPORT.md     人读：可交付报告
    outputs/<project_id>/REPORT.html   可分享：自包含浅色页面
    outputs/<project_id>/report.json   机读：报告结构（供下游消费）
"""

from __future__ import annotations

import html
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from core.config import get_settings
from core.enums import REPORT_FORMATS, ExecStatus, ReportFormat


# 状态 → (中文展示名, 颜色)。颜色仅用于 HTML；通过=绿、失败=红（测试报告惯例）。
_STATUS_STYLE: dict[str, tuple[str, str]] = {
    ExecStatus.PASS.value: ("通过", "#1D9E75"),
    ExecStatus.FAIL.value: ("失败", "#E24B4A"),
    ExecStatus.ERROR.value: ("异常", "#C2410C"),
    ExecStatus.SKIPPED.value: ("跳过", "#BA7517"),
    ExecStatus.STRUCTURAL_ONLY.value: ("仅静态", "#6B7785"),
    ExecStatus.BLOCKED_AUTH.value: ("需鉴权", "#8B5CF6"),
    ExecStatus.BLOCKED_REVIEW.value: ("待审批", "#BA7517"),
    ExecStatus.CANCELLED.value: ("已中断", "#6B7785"),
}

_DIRECTION_LABEL = {
    "improving": "改善 ↗",
    "regressing": "退化 ↘",
    "stable": "平稳 →",
}

_CSS = """
body{font-family:-apple-system,Segoe UI,Roboto,"Microsoft YaHei",sans-serif;margin:32px;
color:#1a1a1a;background:#fff;line-height:1.6;}
h1{font-size:20px;margin:0 0 10px;} h3{margin-top:26px;font-size:15px;
border-left:4px solid #2563eb;padding-left:10px;}
.card{background:#f6f7f9;border:1px solid #e3e6ea;border-radius:10px;padding:14px 16px;margin:12px 0;
font-size:13px;}
.exec{background:#EAF2FF;border:1px solid #B9D4FF;border-left:4px solid #2563EB;border-radius:10px;
padding:14px 16px;margin:12px 0;font-size:13px;}
.exec .headline{font-size:15px;font-weight:600;margin:6px 0;color:#1a1a1a;}
.kpis{display:flex;gap:12px;flex-wrap:wrap;}
.kpi{flex:1;min-width:104px;text-align:center;background:#fff;border:1px solid #e3e6ea;
border-radius:10px;padding:14px;}
.kpi .n{font-size:24px;font-weight:600;}
.kpi .t{font-size:12px;color:#5a6570;}
table{width:100%;border-collapse:collapse;margin-top:8px;}
th,td{text-align:left;padding:7px 9px;border-bottom:1px solid #eee;font-size:12.5px;
vertical-align:top;}
th{background:#f0f2f5;}
code{font-size:11px;color:#555;background:#f0f2f5;padding:1px 4px;border-radius:3px;}
.badge{font-weight:600;}
.muted{color:#5a6570;font-size:12px;}
details{margin-top:6px;} summary{cursor:pointer;font-size:12.5px;}
.ok{color:#1D9E75;}
"""


# ============================================================================
# 通用小工具
# ============================================================================
def _esc(value: Any) -> str:
    return html.escape("" if value is None else str(value))


def _md_escape(value: Any) -> str:
    return str("" if value is None else value).replace("|", "\\|").replace("\n", " ").strip()


def _md_table(header: list[str], rows: list[list[Any]]) -> str:
    out = ["| " + " | ".join(header) + " |", "|" + "|".join(["---"] * len(header)) + "|"]
    for row in rows:
        out.append("| " + " | ".join(_md_escape(c) for c in row) + " |")
    return "\n".join(out)


def _label(status: str) -> str:
    return _STATUS_STYLE.get(status, (status or "—", "#888"))[0]


def _rate_text(rate: float | None) -> str:
    return f"{rate}%" if rate is not None else "—"


def _duration_text(ms: Any) -> str:
    try:
        value = int(ms or 0)
    except (TypeError, ValueError):
        return "—"
    return f"{value / 1000:.2f}s" if value else "—"


@dataclass
class ReportPaths:
    directory: Path
    markdown: Path
    html: Path
    json: Path


class ReportWriter:
    """按项目分目录写出报告三件套。"""

    def __init__(self, base: str | Path | None = None) -> None:
        self.base = Path(base) if base is not None else get_settings().output_dir
        self.base.mkdir(parents=True, exist_ok=True)

    def paths(self, project_id: int) -> ReportPaths:
        d = self.base / str(project_id)
        d.mkdir(parents=True, exist_ok=True)
        return ReportPaths(
            directory=d,
            markdown=d / "REPORT.md",
            html=d / "REPORT.html",
            json=d / "report.json",
        )

    def write(self, project_id: int, report: dict[str, Any]) -> dict[str, str]:
        p = self.paths(project_id)
        p.markdown.write_text(render_markdown(report), encoding="utf-8")
        p.html.write_text(render_html(report), encoding="utf-8")
        p.json.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return {"markdown": str(p.markdown), "html": str(p.html), "json": str(p.json)}


# ============================================================================
# Markdown
# ============================================================================
def _md_header(report: dict[str, Any]) -> list[str]:
    project = report.get("project") or {}
    counts = report.get("counts") or {}
    batch = report.get("batch") or {}
    batch_id = batch.get("batch_id") or "（无执行批次）"
    return [
        f"# 测试报告 · {project.get('name') or '未命名'}（#{project.get('id')}）",
        "",
        f"> 生成时间：{report.get('generated_at', '')}　｜　批次：`{batch_id}`　｜　"
        f"报告版本：{report.get('report_version', '')}　｜　契约版本：{report.get('contract_version', '')}",
        "",
        f"> 规模：用例 {counts.get('cases_active', 0)} 条（在用）/ 测试点 {counts.get('test_points', 0)} 条 / "
        f"功能点 {counts.get('functional_points', 0)} 个",
        "",
        "> 口径：通过率只统计**真实执行并得出结论**的用例（pass / (pass+fail)）；"
        "跳过 / 异常 / 环境态**不计入**通过率，但一律如实列出，避免被误读成「全绿」。",
        "",
    ]


def _md_summary(report: dict[str, Any]) -> list[str]:
    es = report.get("exec_summary") or {}
    metrics = es.get("metrics") or {}
    lines = ["## 一、执行摘要", "", f"**{es.get('headline', '（无结论）')}**", ""]
    lines += [f"- {b}" for b in es.get("bullets") or ["（无）"]]
    lines += [
        "",
        _md_table(
            ["总数", "通过", "失败", "判定类通过率", "跳过", "异常", "需鉴权", "仅静态"],
            [
                [
                    metrics.get("total", 0),
                    metrics.get("pass", 0),
                    metrics.get("fail", 0),
                    _rate_text(metrics.get("pass_rate")),
                    metrics.get("skipped", 0),
                    metrics.get("error", 0),
                    metrics.get("blocked_auth", 0),
                    metrics.get("structural_only", 0),
                ]
            ],
        ),
        "",
    ]
    return lines


def _md_coverage(report: dict[str, Any]) -> list[str]:
    cov = report.get("coverage") or {}
    lines = [
        "## 二、覆盖率（功能点 → 测试点 → 用例）",
        "",
        _md_table(
            ["层级", "总数", "已覆盖", "覆盖率", "缺口"],
            [
                [
                    "功能点 FP",
                    cov.get("fp_total", 0),
                    cov.get("fp_covered", 0),
                    _rate_text(cov.get("fp_rate")),
                    cov.get("uncovered_fp_count", 0),
                ],
                [
                    "测试点 TP",
                    cov.get("tp_total", 0),
                    cov.get("tp_covered", 0),
                    _rate_text(cov.get("tp_rate")),
                    cov.get("uncovered_tp_count", 0),
                ],
                ["用例 Case（在用）", cov.get("case_active", 0), "—", "—", "—"],
            ],
        ),
        "",
        "### 分模块",
        "",
        _md_table(
            ["模块", "功能点", "功能点覆盖", "测试点", "测试点覆盖"],
            [
                [
                    m,
                    d.get("fp_total", 0),
                    f"{d.get('fp_covered', 0)} ({_rate_text(d.get('fp_rate'))})",
                    d.get("tp_total", 0),
                    f"{d.get('tp_covered', 0)} ({_rate_text(d.get('tp_rate'))})",
                ]
                for m, d in (cov.get("module_coverage") or {}).items()
            ]
            or [["（无）", 0, "—", 0, "—"]],
        ),
        "",
    ]
    return lines


def _md_trend(report: dict[str, Any]) -> list[str]:
    tr = report.get("trend") or {}
    lines = [
        f"## 三、执行趋势（{tr.get('count', 0)} 个批次 · 方向："
        f"{_DIRECTION_LABEL.get(str(tr.get('direction')), tr.get('direction', '—'))}）",
        "",
        _md_table(
            ["开始时间", "批次", "终态", "总数", "通过", "失败", "通过率", "跳过", "异常"],
            [
                [
                    p.get("started_at", ""),
                    p.get("batch_id", ""),
                    p.get("state", ""),
                    p.get("total", 0),
                    p.get("pass", 0),
                    p.get("fail", 0),
                    _rate_text(p.get("pass_rate")),
                    p.get("skipped", 0),
                    p.get("error", 0),
                ]
                for p in (tr.get("points") or [])
            ]
            or [["（无执行批次）", "", "", 0, 0, 0, "—", 0, 0]],
        ),
        "",
    ]
    return lines


def _md_flaky(report: dict[str, Any]) -> list[str]:
    fl = report.get("flaky") or {}
    lines = [
        f"## 四、Flaky（不稳定）用例 —— {fl.get('count', 0)} 条",
        "",
        f"> 已评估 {fl.get('evaluated', 0)} 个「跑过多次」的测试点（跨批次 ≥ "
        f"{fl.get('min_batches', 2)} 次才参与判定）。",
        "",
    ]
    if fl.get("count"):
        lines.append(
            _md_table(
                ["测试点", "批次数", "出现过的结论"],
                [
                    [c.get("tp_id"), c.get("batches"), "、".join(c.get("statuses") or [])]
                    for c in fl["cases"]
                ],
            )
        )
    else:
        lines.append("未检测到跨批次结论不一致的用例。")
    lines.append("")
    return lines


def _md_trace(report: dict[str, Any]) -> list[str]:
    tr = report.get("traceability") or {}
    lines = [
        "## 五、追溯体检（结果 → 用例 → 测试点 → 功能点 → 源码）",
        "",
        f"- 功能点 {tr.get('fp_count', 0)} / 测试点 {tr.get('tp_count', 0)} / 用例 {tr.get('case_count', 0)}",
        f"- 孤儿测试点（找不到功能点）：**{tr.get('orphan_tp_count', 0)}**",
        f"- 孤儿用例（找不到测试点）：**{tr.get('orphan_case_count', 0)}**",
        "",
        "> 两个孤儿计数**应恒为 0**；不为 0 说明编号或对账出了问题，需先修数据再谈结论。",
        "",
    ]
    return lines


def _md_evidence(report: dict[str, Any]) -> list[str]:
    evidence = report.get("evidence") or []
    lines = [
        f"## 六、执行证据（逐条，共 {len(evidence)} 条）",
        "",
    ]
    if evidence:
        lines.append(
            _md_table(
                ["case_id", "模块", "用例", "结论", "耗时", "测试点", "功能点", "说明"],
                [
                    [
                        r.get("case_id"),
                        r.get("module") or "—",
                        r.get("title") or "—",
                        _label(str(r.get("status"))),
                        _duration_text(r.get("duration_ms")),
                        r.get("tp_id") or "—",
                        r.get("fp_contract_id") or "—",
                        (r.get("detail") or "")[:120],
                    ]
                    for r in evidence
                ],
            )
        )
    else:
        lines.append("暂无执行证据（本批次未执行或没有单条结论）。")
    lines.append("")
    return lines


def _md_reuse(report: dict[str, Any]) -> list[str]:
    reuse = report.get("reuse") or {}
    lifecycle = reuse.get("lifecycle") or {}
    test_type = reuse.get("test_type") or {}
    return [
        "## 七、用例复用占比",
        "",
        _md_table(
            ["口径", "复用（沿用）", "更新（增量）", "在用合计", "复用占比"],
            [
                [
                    "用例资产（对账后）",
                    reuse.get("reused", 0),
                    reuse.get("updated", 0),
                    reuse.get("active", 0),
                    _rate_text(reuse.get("reuse_rate")),
                ]
            ],
        ),
        "",
        f"> 生命周期分布：{' / '.join(f'{k}×{v}' for k, v in lifecycle.items()) or '（无）'}"
        "（obsolete 不进入执行与导出）。",
        f"> 标签分布：{' / '.join(f'{k}×{v}' for k, v in test_type.items()) or '（无）'}",
        "",
    ]


def _md_change(report: dict[str, Any]) -> list[str]:
    change = report.get("change")
    lines = ["## 八、变更摘要", ""]
    if change:
        lines.append(
            f"- 最近一次对账（`{change.get('kind')}`，{change.get('created_at')}）："
            f"{change.get('detail')}"
        )
    else:
        lines.append("暂无变更对账记录（本服务不做 diff 落库，变更信息来自 change_log）。")
    lines.append("")
    return lines


def render_markdown(report: dict[str, Any]) -> str:
    """渲染人可读报告（Markdown）。"""
    lines: list[str] = []
    for section in (
        _md_header,
        _md_summary,
        _md_coverage,
        _md_trend,
        _md_flaky,
        _md_trace,
        _md_evidence,
        _md_reuse,
        _md_change,
    ):
        lines += section(report)
    return "\n".join(lines).rstrip() + "\n"


# ============================================================================
# HTML
# ============================================================================
def _html_kpis(report: dict[str, Any]) -> str:
    es = report.get("exec_summary") or {}
    metrics = es.get("metrics") or {}
    cov = report.get("coverage") or {}
    cells = [
        (metrics.get("total", 0), "总用例", "#1a1a1a"),
        (metrics.get("pass", 0), "通过", "#1D9E75"),
        (metrics.get("fail", 0), "失败", "#E24B4A"),
        (metrics.get("skipped", 0), "跳过", "#BA7517"),
        (metrics.get("error", 0), "异常", "#C2410C"),
        (_rate_text(metrics.get("pass_rate")), "判定类通过率", "#2563EB"),
        (_rate_text(cov.get("fp_rate")), "功能点覆盖", "#1D9E75"),
        (_rate_text(cov.get("tp_rate")), "测试点覆盖", "#1D9E75"),
    ]
    return (
        "<div class='kpis'>"
        + "".join(
            f"<div class='kpi'><div class='n' style='color:{color}'>{_esc(value)}</div>"
            f"<div class='t'>{_esc(label)}</div></div>"
            for value, label, color in cells
        )
        + "</div>"
    )


def _html_summary(report: dict[str, Any]) -> str:
    es = report.get("exec_summary") or {}
    bullets = "".join(f"<li>{_esc(b)}</li>" for b in (es.get("bullets") or []))
    return (
        "<div class='exec'><b>执行摘要</b>"
        f"<div class='headline'>{_esc(es.get('headline', ''))}</div>"
        f"<ul style='margin:6px 0 2px 18px'>{bullets}</ul></div>"
    )


def _html_coverage(report: dict[str, Any]) -> str:
    cov = report.get("coverage") or {}
    rows = (
        "".join(
            f"<tr><td>{_esc(m)}</td><td>{_esc(d.get('fp_total'))}</td>"
            f"<td>{_esc(d.get('fp_covered'))} ({_esc(_rate_text(d.get('fp_rate')))})</td>"
            f"<td>{_esc(d.get('tp_total'))}</td>"
            f"<td>{_esc(d.get('tp_covered'))} ({_esc(_rate_text(d.get('tp_rate')))})</td></tr>"
            for m, d in (cov.get("module_coverage") or {}).items()
        )
        or "<tr><td colspan='5'>（无模块数据）</td></tr>"
    )
    gaps = (
        f"未覆盖功能点 <b>{cov.get('uncovered_fp_count', 0)}</b> 个　"
        f"未覆盖测试点 <b>{cov.get('uncovered_tp_count', 0)}</b> 个"
    )
    if cov.get("uncovered_fps"):
        items = "".join(
            f"<li><code>{_esc(f.get('contract_id'))}</code> 【{_esc(f.get('ftype'))}】"
            f"{_esc(f.get('name'))}（{_esc(f.get('module')) or '未分类'}）</li>"
            for f in cov["uncovered_fps"]
        )
        gaps += f"<details><summary>未覆盖功能点明细</summary><ul>{items}</ul></details>"
    if cov.get("uncovered_tps"):
        items = "".join(
            f"<li><code>{_esc(t.get('tp_id'))}</code> {_esc(t.get('title'))}"
            f"（{_esc(t.get('module')) or '未分类'}）</li>"
            for t in cov["uncovered_tps"]
        )
        gaps += f"<details><summary>未覆盖测试点明细</summary><ul>{items}</ul></details>"
    return (
        "<h3>覆盖率（功能点 → 测试点 → 用例）</h3>"
        f"<div class='card'>{gaps}</div>"
        "<table><tr><th>模块</th><th>功能点</th><th>功能点覆盖</th>"
        f"<th>测试点</th><th>测试点覆盖</th></tr>{rows}</table>"
    )


def _html_trend(report: dict[str, Any]) -> str:
    tr = report.get("trend") or {}
    rows = (
        "".join(
            f"<tr><td>{_esc(p.get('started_at'))}</td><td><code>{_esc(p.get('batch_id'))}</code></td>"
            f"<td>{_esc(p.get('state'))}</td><td>{_esc(p.get('total'))}</td>"
            f"<td class='ok'>{_esc(p.get('pass'))}</td>"
            f"<td style='color:#E24B4A'>{_esc(p.get('fail'))}</td>"
            f"<td><b>{_esc(_rate_text(p.get('pass_rate')))}</b></td>"
            f"<td>{_esc(p.get('skipped'))}</td><td>{_esc(p.get('error'))}</td></tr>"
            for p in (tr.get("points") or [])
        )
        or "<tr><td colspan='9'>暂无执行批次</td></tr>"
    )
    direction = _DIRECTION_LABEL.get(str(tr.get("direction")), tr.get("direction", "—"))
    return (
        f"<h3>执行趋势（{_esc(tr.get('count', 0))} 个批次 · 方向：{_esc(direction)}）</h3>"
        "<table><tr><th>开始时间</th><th>批次</th><th>终态</th><th>总数</th><th>通过</th>"
        f"<th>失败</th><th>通过率</th><th>跳过</th><th>异常</th></tr>{rows}</table>"
    )


def _html_flaky(report: dict[str, Any]) -> str:
    fl = report.get("flaky") or {}
    if fl.get("count"):
        rows = "".join(
            f"<tr><td><code>{_esc(c.get('tp_id'))}</code></td><td>{_esc(c.get('batches'))}</td>"
            f"<td>{_esc('、'.join(c.get('statuses') or []))}</td></tr>"
            for c in fl["cases"]
        )
        body = f"<table><tr><th>测试点</th><th>批次数</th><th>出现过的结论</th></tr>{rows}</table>"
    else:
        body = (
            f"<div class='card ok'>未检测到跨批次结论不一致的用例"
            f"（已评估 {_esc(fl.get('evaluated', 0))} 个多批次测试点）。</div>"
        )
    return f"<h3>Flaky（不稳定）用例 —— {_esc(fl.get('count', 0))} 条</h3>{body}"


def _html_evidence(report: dict[str, Any]) -> str:
    evidence = report.get("evidence") or []
    if evidence:
        rows = "".join(
            f"<tr><td><code>{_esc(r.get('case_id'))}</code></td><td>{_esc(r.get('module') or '—')}</td>"
            f"<td>{_esc(_label(str(r.get('status'))))}</td>"
            f"<td>{_esc(r.get('title') or '—')}</td>"
            f"<td>{_esc(_duration_text(r.get('duration_ms')))}</td>"
            f"<td><code>{_esc(r.get('tp_id') or '—')}</code></td>"
            f"<td><code>{_esc(r.get('fp_contract_id') or '—')}</code></td>"
            f"<td class='muted'>{_esc((r.get('detail') or '')[:200])}</td></tr>"
            for r in evidence
        )
        body = (
            "<table><tr><th>case_id</th><th>模块</th><th>结论</th><th>用例</th><th>耗时</th>"
            f"<th>测试点</th><th>功能点</th><th>说明</th></tr>{rows}</table>"
        )
    else:
        body = "<div class='card muted'>暂无执行证据（本批次未执行或没有单条结论）。</div>"
    return f"<h3>执行证据（逐条，共 {_esc(len(evidence))} 条）</h3>{body}"


def _html_footer(report: dict[str, Any]) -> str:
    tr = report.get("traceability") or {}
    reuse = report.get("reuse") or {}
    change = report.get("change")
    lifecycle = (
        " / ".join(f"{k}×{v}" for k, v in (reuse.get("lifecycle") or {}).items()) or "（无）"
    )
    change_txt = (
        f"{_esc(change.get('kind'))} @ {_esc(change.get('created_at'))}：{_esc(change.get('detail'))}"
        if change
        else "暂无变更对账记录"
    )
    return (
        "<h3>追溯与复用</h3>"
        "<div class='card'>"
        f"<b>追溯体检：</b>功能点 {_esc(tr.get('fp_count', 0))} / 测试点 {_esc(tr.get('tp_count', 0))} / "
        f"用例 {_esc(tr.get('case_count', 0))}；孤儿测试点 "
        f"<b class='ok'>{_esc(tr.get('orphan_tp_count', 0))}</b>、孤儿用例 "
        f"<b class='ok'>{_esc(tr.get('orphan_case_count', 0))}</b>（应恒为 0）。<br>"
        f"<b>复用占比：</b>{_esc(_rate_text(reuse.get('reuse_rate')))}（复用 {_esc(reuse.get('reused', 0))} / "
        f"更新 {_esc(reuse.get('updated', 0))} / 在用 {_esc(reuse.get('active', 0))}）<br>"
        f"<b>生命周期：</b>{_esc(lifecycle)}<br>"
        f"<b>最近变更：</b>{change_txt}"
        "</div>"
    )


def render_html(report: dict[str, Any]) -> str:
    """渲染自包含浅色 HTML 报告（零外部依赖，可直接分享）。"""
    project = report.get("project") or {}
    title = f"测试报告 · {project.get('name') or '未命名'}（#{project.get('id')}）"
    batch = report.get("batch") or {}
    meta = (
        f"项目：{_esc(project.get('name'))}（#{_esc(project.get('id'))}）　｜　"
        f"生成时间：{_esc(report.get('generated_at'))}　｜　"
        f"批次：<code>{_esc(batch.get('batch_id') or '（无执行批次）')}</code>　｜　"
        f"报告版本：{_esc(report.get('report_version'))}"
    )
    parts = [
        "<!doctype html><html lang='zh'><head><meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width,initial-scale=1'>",
        f"<title>{_esc(title)}</title>",
        f"<style>{_CSS}</style>",
        "</head><body>",
        f"<h1>{_esc(title)}</h1>",
        f"<div class='card'>{meta}</div>",
        "<div class='card muted'>口径：通过率只统计<strong>真实执行并得出结论</strong>的用例"
        "（pass / (pass+fail)）；跳过 / 异常 / 环境态<strong>不计入</strong>通过率，但一律如实列出，"
        "避免被误读成「全绿」。</div>",
        _html_summary(report),
        _html_kpis(report),
        _html_coverage(report),
        _html_trend(report),
        _html_flaky(report),
        _html_evidence(report),
        _html_footer(report),
        "</body></html>",
    ]
    return "\n".join(parts) + "\n"


# ============================================================================
# 可选导出格式（G-11）：PDF / Word / Excel —— 依赖按需引入，缺依赖给友好提示
# ============================================================================
class ReportExportError(ValueError):
    """导出失败：缺失可选依赖（reportlab / python-docx / openpyxl）或非受支持格式。

    继承自 ValueError（而非 RuntimeError），便于 HTTP 层统一转 422 并附 `pip install` 提示。
    """


def export_report(report: dict[str, Any], fmt: str, path: str) -> str:
    """把报告结构导出为指定格式落盘，返回写出路径。

    - `json` / `md` / `html`：内置渲染器，无需额外依赖；
    - `pdf` / `docx` / `xlsx`：按需 `importlib` 引入可选依赖，**缺依赖则抛
      `ReportExportError`（附 pip install 命令）**，绝不静默产出半截文件。

    所有格式均来自同一份 `report` 字典（DB 事实的纯函数），口径与 `engine/report.py`
    完全一致。凭证 / 令牌不进入报告内容，故导出物天然不含敏感字段。
    """
    from core.enums import REPORT_FORMAT_OPTIONAL_DEPS

    fmt = (fmt or "").lower().strip()
    if fmt not in REPORT_FORMATS:
        raise ReportExportError(f"不支持的报告格式：{fmt!r}（允许 {sorted(REPORT_FORMATS)}）")

    if fmt in (ReportFormat.JSON.value, ReportFormat.MARKDOWN.value, ReportFormat.HTML.value):
        if fmt == ReportFormat.JSON.value:
            rendered = json.dumps(report, ensure_ascii=False, indent=2)
        elif fmt == ReportFormat.MARKDOWN.value:
            rendered = render_markdown(report)
        else:
            rendered = render_html(report)
        Path(path).write_text(rendered + "\n", encoding="utf-8")
        return path

    # 可选格式：缺依赖直接报错（不再尝试回退，避免误导用户以为导出成功）
    dep = REPORT_FORMAT_OPTIONAL_DEPS[fmt]
    try:
        if fmt == ReportFormat.PDF.value:
            _render_pdf(report, path)
        elif fmt == ReportFormat.WORD.value:
            _render_docx(report, path)
        else:  # xlsx
            _render_xlsx(report, path)
    except ImportError as exc:  # 缺依赖：明确提示
        raise ReportExportError(
            f"导出 {fmt} 需要安装可选依赖 `{dep}`（pip install {dep}）；"
            f"当前环境未安装，无法导出：{exc}"
        ) from exc
    return path


def _render_pdf(report: dict[str, Any], path: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    body = getSampleStyleSheet()["BodyText"]
    h1 = ParagraphStyle("h1", parent=body, fontSize=15, spaceAfter=8)
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm)
    story: list[Any] = [Paragraph(_plain_title(report), h1), Spacer(1, 6)]
    story.append(Paragraph(_plain_summary(report).replace("\n", "<br/>"), body))
    for line in _plain_lines(report):
        story.append(Paragraph(_esc(line), body))
    doc.build(story)


def _render_docx(report: dict[str, Any], path: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(_plain_title(report), level=1)
    doc.add_paragraph(_plain_summary(report))
    for line in _plain_lines(report):
        doc.add_paragraph(line)
    doc.save(path)


def _render_xlsx(report: dict[str, Any], path: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "测试报告"
    ws.append([_plain_title(report)])
    ws.append([_plain_summary(report)])
    ws.append([])
    for line in _plain_lines(report):
        ws.append([line])
    wb.save(path)


def _plain_title(report: dict[str, Any]) -> str:
    project = report.get("project") or {}
    return f"测试报告 · {project.get('name') or '未命名'}（#{project.get('id')}）"


def _plain_summary(report: dict[str, Any]) -> str:
    generated = report.get("generated_at", "")
    batch = (report.get("batch") or {}).get("batch_id") or "（无执行批次）"
    return f"生成时间：{generated}　批次：{batch}　报告版本：{report.get('report_version')}"


def _plain_lines(report: dict[str, Any]) -> list[str]:
    """把报告中可结构化的部分铺平成纯文本行（用于 pdf/docx/xlsx 同源渲染）。"""
    lines: list[str] = []
    summary = report.get("summary") or {}
    if summary:
        lines.append("## 概要")
        for k, v in summary.items():
            lines.append(f"- {k}: {v}")
    coverage = report.get("coverage") or {}
    if coverage:
        lines.append("## 覆盖率")
        for k, v in coverage.items():
            lines.append(f"- {k}: {v}")
    cases = report.get("cases") or []
    if cases:
        lines.append("## 用例清单")
        for c in cases[:500]:  # 防 XLSX 单行数爆炸；完整清单以 DB 为准
            lines.append(
                f"- [{c.get('case_type')}] {c.get('title')} → {c.get('exec_status') or '未执行'}"
            )
    if not lines:
        lines.append("（报告无可结构化条目）")
    return lines
